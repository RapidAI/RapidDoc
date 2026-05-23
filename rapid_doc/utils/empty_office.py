from functools import lru_cache
from io import BytesIO


EMPTY_OFFICE_SUFFIXES = {"docx", "pptx", "xlsx", "xlsm"}


def normalize_empty_office_bytes(file_bytes: bytes, file_suffix: str | None) -> bytes:
    suffix = (file_suffix or "").lower().lstrip(".")
    if file_bytes or suffix not in EMPTY_OFFICE_SUFFIXES:
        return file_bytes
    return _empty_office_template(suffix)


@lru_cache(maxsize=len(EMPTY_OFFICE_SUFFIXES))
def _empty_office_template(file_suffix: str) -> bytes:
    output = BytesIO()
    if file_suffix == "docx":
        from docx import Document

        Document().save(output)
    elif file_suffix == "pptx":
        from pptx import Presentation

        Presentation().save(output)
    else:
        from openpyxl import Workbook

        Workbook().save(output)
    return output.getvalue()
